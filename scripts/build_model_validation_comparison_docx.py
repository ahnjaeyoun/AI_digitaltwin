"""모델 학습 방식 비교와 5겹 검증 선정 이유를 정리한 DOCX를 생성합니다.

이 문서는 발표 자료가 아니라, 개발 방향을 빠르게 확인하기 위한 1페이지 요약 문서입니다.
기존 DOCX는 수정하지 않고 새로운 파일만 생성합니다.
"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_ROOT / "DigitalTwin_모델학습방식_비교와_5겹검증_선정.docx"

# compact_reference_guide 프리셋에서 사용하는 기본 색상과 표 크기입니다.
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
BODY_COLOR = "1F1F1F"
MUTED_COLOR = "666666"
TABLE_HEADER_FILL = "E8EEF5"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(
    run,
    size: float,
    *,
    bold: bool = False,
    color: str = BODY_COLOR,
) -> None:
    """한글과 영문이 Word에서 같은 크기로 보이도록 글꼴을 명시합니다."""

    run.font.name = "맑은 고딕"
    run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    """표 머리글 셀에 옅은 배경색을 적용합니다."""

    cell_properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell_properties.append(shading)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    """셀 안쪽 여백을 명시하여 글자가 표 테두리에 붙지 않게 합니다."""

    cell_properties = cell._tc.get_or_add_tcPr()
    margins = cell_properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        cell_properties.append(margins)

    for side, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        margin = margins.find(qn(f"w:{side}"))
        if margin is None:
            margin = OxmlElement(f"w:{side}")
            margins.append(margin)
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")


def set_fixed_table_geometry(table, column_widths: list[int]) -> None:
    """Word와 PDF에서 표 폭이 달라지지 않도록 정확한 DXA 폭을 설정합니다."""

    if sum(column_widths) != TABLE_WIDTH_DXA:
        raise ValueError("표 열 너비의 합은 9360 DXA여야 합니다.")

    table.autofit = False
    table_properties = table._tbl.tblPr

    table_width = table_properties.first_child_found_in("w:tblW")
    table_width.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    table_width.set(qn("w:type"), "dxa")

    table_indent = table_properties.first_child_found_in("w:tblInd")
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        table_properties.append(table_indent)
    table_indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    table_indent.set(qn("w:type"), "dxa")

    for grid_column, width in zip(table._tbl.tblGrid.gridCol_lst, column_widths):
        grid_column.set(qn("w:w"), str(width))

    for row in table.rows:
        for cell, width in zip(row.cells, column_widths):
            cell_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_heading(document: Document, text: str) -> None:
    """문서 안의 짧은 구역 제목을 추가합니다."""

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.keep_with_next = True
    set_run_font(paragraph.add_run(text), 13, bold=True, color=BLUE)


def add_label_paragraph(document: Document, label: str, description: str) -> None:
    """추천 이유를 '짧은 제목 + 쉬운 설명' 형식으로 추가합니다."""

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.15
    set_run_font(paragraph.add_run(f"{label}: "), 10.2, bold=True, color=DARK_BLUE)
    set_run_font(paragraph.add_run(description), 10.2)


def add_comparison_table(document: Document) -> None:
    """세 가지 모델 학습 방식을 같은 기준으로 비교하는 표를 만듭니다."""

    headers = ["방법", "장점", "단점"]
    rows = [
        [
            "1. LightGBM\n+ 5겹 그룹 검증",
            "표 형태 센서 데이터에 적합하고, 5번 검증해 결과가 안정적입니다. 실시간 예측도 빠릅니다.",
            "학습을 여러 번 수행하므로 단일 분할보다 학습 시간이 더 필요합니다.",
        ],
        [
            "2. LightGBM\n+ 단일 80:20 분할",
            "학습이 가장 빠르고 코드가 단순합니다.",
            "한 번의 데이터 분할에 따라 평가 점수가 우연히 달라질 수 있습니다.",
        ],
        [
            "3. LSTM\n시계열 신경망",
            "이전 센서값의 시간 흐름을 직접 학습할 수 있습니다.",
            "전처리·학습·설명이 복잡하고 Unity 실시간 연동 부담이 커집니다.",
        ],
    ]

    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"

    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, TABLE_HEADER_FILL)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        set_run_font(paragraph.add_run(header), 9.6, bold=True, color=DARK_BLUE)

    for values in rows:
        cells = table.add_row().cells
        for column_index, (cell, value) in enumerate(zip(cells, values)):
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            if column_index == 0:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(
                paragraph.add_run(value),
                9.2,
                bold=(column_index == 0),
                color=DARK_BLUE if column_index == 0 else BODY_COLOR,
            )

    set_fixed_table_geometry(table, [2100, 3730, 3530])


def configure_document_styles(document: Document) -> None:
    """페이지와 본문 스타일을 compact_reference_guide 값으로 설정합니다."""

    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "맑은 고딕"
    normal_style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal_style.font.size = Pt(11)
    normal_style.paragraph_format.space_after = Pt(6)
    normal_style.paragraph_format.line_spacing = 1.25


def build_document(output_path: Path = OUTPUT_PATH) -> Path:
    """학습 방식 비교와 최종 선정 내용을 담은 새 DOCX를 생성합니다."""

    document = Document()
    configure_document_styles(document)

    # memo_masthead 패턴을 간단히 적용한 제목 영역입니다.
    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    set_run_font(
        title.add_run("DigitalTwin 모델 학습 방식 비교"),
        20,
        bold=True,
        color=DARK_BLUE,
    )

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(9)
    set_run_font(
        subtitle.add_run("5겹 그룹 검증 선정 이유 | 2026-07-21"),
        10,
        color=MUTED_COLOR,
    )

    add_heading(document, "1. 최종 결정")
    decision = document.add_paragraph()
    decision.paragraph_format.space_after = Pt(5)
    decision.paragraph_format.line_spacing = 1.15
    set_run_font(decision.add_run("선택: "), 10.5, bold=True, color=DARK_BLUE)
    set_run_font(
        decision.add_run(
            "LightGBM + StratifiedGroupKFold 5겹 검증. "
            "완성 데이터셋(final_training_dataset_solver_latest.csv)으로 매초 실시간 판정 모델을 학습합니다."
        ),
        10.5,
    )

    add_heading(document, "2. 세 가지 방법 비교")
    add_comparison_table(document)

    add_heading(document, "3. 1번을 추천한 이유")
    add_label_paragraph(
        document,
        "데이터 적합성",
        "압력·유량·온도·밸브 위치처럼 한 행에 정리된 센서값은 LightGBM과 잘 맞습니다.",
    )
    add_label_paragraph(
        document,
        "누출 방지",
        "같은 cycle_id의 여러 초 데이터가 학습과 검증에 나뉘지 않도록 한 묶음으로 처리합니다.",
    )
    add_label_paragraph(
        document,
        "평가 신뢰성",
        "서로 다른 검증 묶음으로 5번 평가하고 평균과 편차를 저장하므로 단일 분할보다 결과를 믿기 쉽습니다.",
    )
    add_label_paragraph(
        document,
        "실시간 사용",
        "5번 반복은 학습할 때만 필요하며, 저장된 최종 모델의 매초 예측 속도는 빠릅니다.",
    )

    add_heading(document, "4. 적용 대상")
    add_label_paragraph(
        document,
        "모델 A",
        "정상(none)·초기(initial)·중기(moderate)·심각(severe) 4단계 판정",
    )
    add_label_paragraph(
        document,
        "모델 B",
        "위험 부품 계통 판정. 외부 누유는 위치를 나누지 않고 배관·연결부 한 그룹으로 처리",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


if __name__ == "__main__":
    print(build_document())
