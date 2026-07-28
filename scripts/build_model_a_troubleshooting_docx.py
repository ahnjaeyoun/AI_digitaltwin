"""DigitalTwin 모델 A 관련 트러블슈팅과 결정 기록 DOCX를 생성합니다.

이 문서는 발표용이 아니라, 데이터·솔버·라벨·검증 문제를 나중에 다시 추적하기 위한
간결한 기술 기록입니다. 기존 DOCX는 수정하지 않고 새 문서만 생성합니다.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

try:
    # `python -m scripts.build_model_a_troubleshooting_docx` 또는 테스트 import 때 사용합니다.
    from .build_model_a_summary_docx import (
        DARK_BLUE,
        HEADER_FILL,
        TEXT,
        add_label_paragraph,
        configure_document,
        set_cell_margins,
        set_cell_shading,
        set_repeat_table_header,
        set_run_font,
        set_table_geometry,
    )
except ImportError:
    # `python scripts/build_model_a_troubleshooting_docx.py`로 직접 실행할 때 사용합니다.
    from build_model_a_summary_docx import (
        DARK_BLUE,
        HEADER_FILL,
        TEXT,
        add_label_paragraph,
        configure_document,
        set_cell_margins,
        set_cell_shading,
        set_repeat_table_header,
        set_run_font,
        set_table_geometry,
    )


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_ROOT / "DigitalTwin_모델A_트러블슈팅_및_결정기록.docx"


def add_table(
    document: Document,
    headers: list[str],
    rows: list[tuple[str, ...]],
    widths: list[int],
    *,
    center_columns: set[int] | None = None,
) -> None:
    """짧은 비교·조회 정보를 고정 폭 표로 추가합니다."""

    center_columns = center_columns or set()
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])

    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, HEADER_FILL)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        set_run_font(paragraph.add_run(header), 9.3, bold=True, color=DARK_BLUE)

    for values in rows:
        cells = table.add_row().cells
        for column_index, (cell, value) in enumerate(zip(cells, values)):
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.08
            if column_index in center_columns:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(
                paragraph.add_run(value),
                9.1,
                bold=(column_index == 0),
                color=DARK_BLUE if column_index == 0 else TEXT,
            )
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    set_table_geometry(table, widths)


def add_issue(
    document: Document,
    title: str,
    *,
    symptom: str,
    cause: str,
    check: str,
    solution: str,
    prevention: str,
) -> None:
    """하나의 문제를 현상·원인·확인·해결·재발 방지 순으로 기록합니다."""

    document.add_heading(title, level=2)
    add_label_paragraph(document, "현상", symptom)
    add_label_paragraph(document, "원인", cause)
    add_label_paragraph(document, "확인", check)
    add_label_paragraph(document, "해결", solution)
    add_label_paragraph(document, "재발 방지", prevention)


def build_document(output_path: Path = OUTPUT_PATH) -> Path:
    """대화에서 확인된 주요 문제와 확정 결정을 기술 기록으로 만듭니다."""

    document = Document()
    configure_document(document)

    # 기술 기록에 적합한 memo_masthead 형식의 간단한 제목 영역입니다.
    document.add_paragraph("DigitalTwin 모델 A 트러블슈팅", style="Title")
    document.add_paragraph("데이터·솔버·라벨·검증 결정 기록 | 2026-07-21", style="Subtitle")

    document.add_heading("1. 현재 확정 상태", level=1)
    add_table(
        document,
        ["항목", "확정 내용"],
        [
            ("현재 모델", "모델 A: 매초 센서값으로 정상·초기·중기·심각 4단계 분류"),
            ("정답", "severity_level의 none·initial·moderate·severe"),
            ("데이터", "dataset/final_training_dataset_solver_latest.csv"),
            ("학습·검증", "LightGBM + cycle_id 기준 StratifiedGroupKFold 5겹 검증"),
            ("보류", "모델 B 위험 부품 판정, Unity 연동, 실시간 예측 프로그램"),
            ("제외", "고장 시점, RUL, 열화 속도, 고장 단계, 고장유형 분류"),
        ],
        [2700, 6660],
        center_columns={0},
    )

    document.add_heading("2. 방향이 변경된 과정", level=1)
    add_table(
        document,
        ["단계", "검토 내용", "최종 판단"],
        [
            ("초기", "고장유형·건전도·RUL을 모두 학습", "RUL에 필요한 연속 수명 이력이 없어 제거"),
            ("예지보전 재정의", "고장 시점, RUL, 열화 속도, 위험 부품, 단계", "현재 데이터로 확실한 것은 현재 열화 단계와 제한적 부품 분류"),
            ("역할 분리", "고장 여부와 고장유형", "다른 팀원이 담당하므로 모델 A에서 제외"),
            ("모델 A", "정상·초기·중기·심각·고장", "고장 단계를 빼고 4단계로 확정"),
            ("구현 순서", "모델 A와 B 동시 개발", "모델 A를 먼저 Python에서 검증하고 모델 B는 추가 논의"),
        ],
        [1500, 3800, 4060],
        center_columns={0},
    )

    document.add_heading("3. 데이터 생성 흐름", level=1)
    add_table(
        document,
        ["파일", "역할", "주의점"],
        [
            ("scenarios.csv", "운전조건과 고장 시나리오·라벨", "gen_scenarios.py가 만든 기존 입력이며 Codex가 새로 만든 파일이 아님"),
            ("gen_output_latest_utf8.csv", "scenarios.csv를 최신 솔버에 넣은 원시 출력", "N_CYL_CAP_p 같은 솔버 컬럼명을 사용"),
            ("final_training_dataset_solver_latest.csv", "시나리오 라벨과 솔버 출력을 합치고 학습용 컬럼명으로 정리한 완성본", "모델 A의 유일한 원본 데이터"),
        ],
        [3100, 3260, 3000],
    )
    add_label_paragraph(
        document,
        "완성본 확인",
        "1,450,300행·100,000사이클·48컬럼이며 정상 30,000, 초기 23,258, 중기 23,238, 심각 23,504사이클",
    )
    add_label_paragraph(
        document,
        "행 단위 일치",
        "최신 솔버 출력과 완성 학습데이터의 행 ID·라벨·센서 매핑을 전수 비교했고 의미 있는 불일치는 없었음",
    )

    document.add_heading("4. 해결한 주요 문제", level=1)

    add_issue(
        document,
        "4.1 잘못된 latest 파일 선택",
        symptom="final_training_dataset_latest.csv를 학습에 사용하려 했으나 4단계 데이터가 보이지 않음",
        cause="이 파일은 생성 도중 끝난 중간 파일로 400,000행·27,569사이클이 모두 normal/none이었음",
        check="행 수, 고유 cycle_id 수, severity_level과 fault_class 분포를 확인",
        solution="완성본 final_training_dataset_solver_latest.csv로 변경",
        prevention="config.py의 기본 경로를 완성본으로 고정하고 preparation.py가 네 단계 존재 여부를 검사",
    )

    add_issue(
        document,
        "4.2 v3와 최신 솔버 압력 불일치",
        symptom="같은 조건인데 v3의 cyl_cap_pressure_bar와 최신 N_CYL_CAP_p가 다르게 보임",
        cause="final_training_dataset_v3.csv는 현재 최신 솔버 결과와 일부 실린더 압력이 맞지 않았고 ext_leak 90,945행에서 의미 있는 차이가 확인됨",
        check="예: row_id 78419_15에서 v3 CAP 54.773 bar, 최신 솔버 CAP 29.94690 bar",
        solution="최신 솔버 출력을 기준으로 final_training_dataset_solver_latest.csv를 다시 구성하고 전수 비교",
        prevention="데이터 버전마다 시나리오 입력·솔버 실행 파일·원시 출력·최종 학습파일을 함께 기록",
    )

    add_issue(
        document,
        "4.3 ext_leak와 솔버 출력의 의미 혼동",
        symptom="ext_leak가 솔버가 계산해 낸 숫자 컬럼인지, 외부 누유도 솔버 출력인지 혼동",
        cause="고장 조건 라벨과 물리 계산 결과가 같은 CSV에 합쳐져 있었음",
        check="ext_leak는 scenarios.csv의 fault_class 값이고 N_CYL_CAP_p는 원시 솔버 압력 출력임을 컬럼 출처로 구분",
        solution="ext_leak는 주입한 고장 시나리오 정답, N_CYL_CAP_p→cyl_cap_pressure_bar는 그 조건에서 계산된 센서값으로 정의",
        prevention="라벨·솔버 입력·솔버 출력·학습 입력을 문서와 config에서 별도 목록으로 관리",
    )

    add_issue(
        document,
        "4.4 solver_warning 계열이 모델 입력처럼 보인 문제",
        symptom="solver_warning, contains_inf_or_nan, exit_code가 모델 학습에 영향을 줄 수 있다고 판단",
        cause="솔버 실행 상태를 확인하는 감사 컬럼과 실제 센서 컬럼의 역할이 구분되지 않았음",
        check="최신 원시 출력에서 solver_warning=no, contains_inf_or_nan=no, status=ok이고 최종 48컬럼 학습 입력에는 포함되지 않음을 확인",
        solution="세 값은 센서가 아닌 검증 조건으로만 사용하고 이상 행은 학습 전에 차단",
        prevention="SENSOR_FEATURES 허용 목록 방식으로 학습 입력을 지정해 예상하지 않은 컬럼이 자동 유입되지 않게 함",
    )

    add_issue(
        document,
        "4.5 존재하지 않는 센서로 고장유형을 설명한 문제",
        symptom="펌프 진동·탱크 레벨·실린더 누설량·밸브 지연 등 실제 데이터에 없던 값이 고장 근거 문서에 등장",
        cause="물리적으로 있으면 좋은 센서 목록과 당시 CSV에 실제 존재하는 컬럼이 섞여 설명됨",
        check="CSV 헤더를 기준으로 실제 측정·계산 컬럼과 추가가 필요한 센서를 다시 대조",
        solution="고장유형 설명 모델을 현재 모델 A에서 제외하고, 모델 A는 실제 존재하는 26개 운전값만 사용",
        prevention="문서에서 실제 컬럼, 계산 가능 컬럼, 향후 추가 센서를 반드시 별도 구역으로 표시",
    )

    add_issue(
        document,
        "4.6 전처리 CSV를 다시 만들어야 한다는 혼동",
        symptom="preparation.py가 새로운 대용량 학습 CSV를 만드는 것으로 이해",
        cause="데이터 준비와 데이터 생성을 같은 작업으로 생각함",
        check="현재 원본 완성본이 이미 존재하고 모델 A에는 단순 선택·검사만 필요함을 확인",
        solution="preparation.py는 원본을 읽고 메모리에서 검사·선택한 뒤 DataFrame을 train.py로 전달",
        prevention="새 CSV 대신 data_validation_result.json만 저장하고 원본 데이터 경로를 config.py 한곳에서 관리",
    )

    add_issue(
        document,
        "4.7 평가 파일이 정답을 어떻게 아는지 혼동",
        symptom="evaluation.py에 정답을 별도로 입력하지 않았는데 채점이 가능한지 의문",
        cause="CSV의 severity_level이 이미 시나리오 정답이라는 점이 명확하지 않았음",
        check="모델 입력에서는 severity_level을 제외하고, 예측 후 evaluation.py만 원래 라벨과 비교",
        solution="입력 X=26개 운전값, 정답 y=severity_level, 그룹=cycle_id로 명확히 분리",
        prevention="결과 JSON에 target_column과 validation_scope를 함께 기록",
    )

    document.add_heading("5. 모델 A 입력과 제외값", level=1)
    add_table(
        document,
        ["구분", "컬럼"],
        [
            ("운전 문맥", "cycle_second, cycle_phase, target_chamber"),
            ("온도·펌프", "temperature_c, pump_rpm"),
            ("밸브", "valve_cmd_pa, valve_pos_pa, valve_cmd_pb, valve_pos_pb, valve_cmd_at, valve_pos_at, valve_cmd_bt, valve_pos_bt"),
            ("유량", "flow_rate_L_min, return_flow_L_min"),
            ("압력", "pump_in_pressure_bar, pump_out_pressure_bar, cyl_cap_pressure_bar, cyl_rod_pressure_bar, pump_delta_pressure_bar"),
            ("계산·설정", "pump_head_m, hydraulic_power_kW, max_velocity_m_s, target_pressure_error_bar, target_pressure_bar, relief_set_pressure_bar"),
            ("입력 제외", "모든 gt_*, fault_class, severity_level, fault_severity, onset_type, is_anomaly, cycle_id"),
        ],
        [2200, 7160],
        center_columns={0},
    )

    document.add_heading("6. 학습 방식 결정", level=1)
    add_table(
        document,
        ["방법", "장점", "판단"],
        [
            ("LightGBM + 5겹 그룹 검증", "표 센서 데이터에 적합, 검증 안정성, 빠른 추론", "채택"),
            ("LightGBM + 단일 80:20", "가장 빠르고 단순", "한 번의 분할 결과에 흔들려 보류"),
            ("LSTM", "시간 흐름 직접 학습 가능", "복잡도·설명·연동 부담으로 보류"),
        ],
        [2700, 4360, 2300],
        center_columns={0, 2},
    )
    add_label_paragraph(
        document,
        "핵심 규칙",
        "같은 cycle_id의 모든 초를 한 그룹으로 유지해 학습과 검증에 동시에 들어가지 않게 함",
    )
    add_label_paragraph(
        document,
        "실시간 의미",
        "학습은 5번 반복하지만 저장 모델의 매초 한 행 예측은 빠름. 사이클 초반 정보 부족은 초별 정확도로 별도 확인",
    )

    document.add_heading("7. 결과값과 검증 범위", level=1)
    add_table(
        document,
        ["결과", "의미"],
        [
            ("4단계 판정", "none=정상, initial=초기, moderate=중기, severe=심각"),
            ("단계별 확률", "네 단계 각각의 예측 확률과 가장 큰 확률인 신뢰도"),
            ("accuracy", "전체 검증 행 중 단계를 맞힌 비율"),
            ("macro F1", "네 단계를 같은 비중으로 평가한 F1 평균"),
            ("단계별 precision·recall·F1", "단계별 오탐과 미탐을 확인"),
            ("혼동행렬", "실제 단계가 어떤 단계로 잘못 예측됐는지 확인"),
            ("5겹 평균·표준편차", "평균 성능과 데이터 분할에 따른 흔들림 확인"),
            ("cycle_second별 정확도", "매초 판정에서 특히 사이클 초반이 약한지 확인"),
        ],
        [3000, 6360],
    )
    add_label_paragraph(
        document,
        "말할 수 있는 것",
        "솔버·시나리오 합성 데이터의 보지 않은 사이클에서 네 단계가 얼마나 구분되는지",
    )
    add_label_paragraph(
        document,
        "말할 수 없는 것",
        "실제 공장에서 동일한 정확도로 동작한다는 주장. 실제 센서 이력과 점검으로 확정된 단계 데이터가 필요",
    )

    document.add_heading("8. 향후 실행 중 빠른 점검표", level=1)
    add_table(
        document,
        ["현상", "먼저 확인", "조치"],
        [
            ("네 단계가 안 보임", "DATASET_PATH와 severity_level 분포", "solver_latest 완성본 사용"),
            ("정확도가 비정상적으로 높음", "gt_*·fault_severity 유입, cycle_id 중복 분리", "허용 입력 목록과 그룹 겹침 검사"),
            ("정확도가 낮음", "단계별 recall, 혼동행렬, cycle_second별 정확도", "약한 단계·초 구간과 생성 규칙 점검"),
            ("필수 컬럼 오류", "CSV 헤더와 config SENSOR_FEATURES", "컬럼명을 맞추거나 config 수정 후 재학습"),
            ("결측·무한대 오류", "data_validation_result.json", "원본 행과 솔버 상태를 확인하고 학습 중단"),
            ("5겹 중 단계 누락", "각 fold의 라벨 분포", "StratifiedGroupKFold 설정과 그룹 라벨 일관성 점검"),
            ("메모리 부족", "읽는 컬럼 수와 dtype", "usecols 사용, float32·category 적용"),
            ("실제 데이터 성능 저하", "합성·실제 센서 범위와 단위", "실제 라벨 데이터로 외부 검증·재학습"),
        ],
        [2700, 3360, 3300],
    )

    document.add_heading("9. 현재 남은 작업", level=1)
    add_table(
        document,
        ["순서", "작업"],
        [
            ("1", "모델 A config.py, preparation.py, train.py, evaluation.py와 자동 테스트 구현"),
            ("2", "필요 패키지 설치 후 5겹 검증과 최종 학습 실행"),
            ("3", "data_validation_result.json과 evaluation_results.json 검토"),
            ("4", "모델 B의 부품 그룹과 정답 정의 추가 논의"),
            ("5", "실제 공정 데이터 확보 시 외부 검증"),
        ],
        [1400, 7960],
        center_columns={0},
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


if __name__ == "__main__":
    print(build_document())
