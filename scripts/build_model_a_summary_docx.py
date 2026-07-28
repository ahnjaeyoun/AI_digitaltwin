"""모델 A 설계와 검증 범위를 짧게 정리한 DOCX를 생성합니다.

발표용 자료가 아니라 개발 방향을 합의하고 나중에 다시 확인하기 위한 간단한 문서입니다.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_ROOT / "DigitalTwin_모델A_열화단계분류_간단정리.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
TEXT = "1F1F1F"
MUTED = "666666"
HEADER_FILL = "E8EEF5"
TABLE_WIDTH_DXA = 9360


def set_run_font(run, size: float, *, bold: bool = False, color: str = TEXT) -> None:
    """Word에서 한글과 영문 글꼴이 안정적으로 표시되도록 설정합니다."""

    run.font.name = "맑은 고딕"
    run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def configure_style(style, *, size: float, color: str, before: float, after: float, line: float) -> None:
    """문단 스타일에 compact_reference_guide의 수치를 명시합니다."""

    style.font.name = "맑은 고딕"
    style._element.get_or_add_rPr()
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line


def configure_document(document: Document) -> None:
    """용지, 여백, 본문과 제목 스타일을 한 번에 설정합니다."""

    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    configure_style(document.styles["Normal"], size=11, color=TEXT, before=0, after=6, line=1.25)
    configure_style(document.styles["Title"], size=21, color=DARK_BLUE, before=0, after=4, line=1.0)
    document.styles["Title"].font.bold = True
    configure_style(document.styles["Subtitle"], size=10, color=MUTED, before=0, after=10, line=1.0)
    configure_style(document.styles["Heading 1"], size=16, color=BLUE, before=18, after=10, line=1.0)
    document.styles["Heading 1"].font.bold = True
    configure_style(document.styles["Heading 2"], size=13, color=BLUE, before=14, after=7, line=1.0)
    document.styles["Heading 2"].font.bold = True


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)

    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        margin = margins.find(qn(f"w:{side}"))
        if margin is None:
            margin = OxmlElement(f"w:{side}")
            margins.append(margin)
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    """표 전체와 각 셀의 폭을 같은 DXA 값으로 고정합니다."""

    if sum(widths) != TABLE_WIDTH_DXA:
        raise ValueError("표 열 너비의 합은 9360 DXA여야 합니다.")

    table.autofit = False
    properties = table._tbl.tblPr
    table_width = properties.first_child_found_in("w:tblW")
    table_width.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    table_width.set(qn("w:type"), "dxa")

    table_indent = properties.first_child_found_in("w:tblInd")
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        properties.append(table_indent)
    table_indent.set(qn("w:w"), "120")
    table_indent.set(qn("w:type"), "dxa")

    for grid_column, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid_column.set(qn("w:w"), str(width))

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row) -> None:
    """표가 다음 페이지로 넘어가도 첫 번째 머리글 행이 반복되게 합니다."""

    row_properties = row._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    row_properties.append(repeat_header)


def add_label_detail_table(document: Document, headers: list[str], rows: list[tuple[str, str]]) -> None:
    """짧은 항목과 설명을 두 열로 비교하기 좋은 표로 만듭니다."""

    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, HEADER_FILL)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        set_run_font(paragraph.add_run(header), 9.5, bold=True, color=DARK_BLUE)

    for label, detail in rows:
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, (label, detail))):
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.1
            if index == 0:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(
                paragraph.add_run(value),
                9.5,
                bold=(index == 0),
                color=DARK_BLUE if index == 0 else TEXT,
            )

    set_table_geometry(table, [2700, 6660])


def add_input_table(document: Document, rows: list[tuple[str, str, str]]) -> None:
    """모델 입력 컬럼을 구분·컬럼명·설명 세 열로 자세히 나열합니다."""

    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for cell, header in zip(table.rows[0].cells, ("구분", "컬럼명", "설명")):
        set_cell_shading(cell, HEADER_FILL)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        set_run_font(paragraph.add_run(header), 9.2, bold=True, color=DARK_BLUE)

    for category, column_name, description in rows:
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, (category, column_name, description))):
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            if index == 0:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(
                paragraph.add_run(value),
                8.8,
                bold=(index == 1),
                color=DARK_BLUE if index == 1 else TEXT,
            )

    set_table_geometry(table, [1500, 3100, 4760])


def add_metric_table(document: Document, rows: list[tuple[str, str]]) -> None:
    """평가 결과 이름과 초보자용 의미를 두 열로 설명합니다."""

    add_label_detail_table(document, ["결과값", "의미"], rows)


def add_label_paragraph(document: Document, label: str, detail: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.2
    set_run_font(paragraph.add_run(f"{label}: "), 10.5, bold=True, color=DARK_BLUE)
    set_run_font(paragraph.add_run(detail), 10.5)


def build_document(output_path: Path = OUTPUT_PATH) -> Path:
    """지금까지 합의한 모델 A의 범위와 검증 의미를 짧게 정리합니다."""

    document = Document()
    configure_document(document)

    # 기술 메모에 맞는 memo_masthead 패턴을 간단히 적용합니다.
    document.add_paragraph("DigitalTwin 모델 A 간단 정리", style="Title")
    document.add_paragraph("열화 단계 분류 설계와 검증 범위 | 2026-07-21", style="Subtitle")

    document.add_heading("1. 이번에 만드는 것", level=1)
    add_label_detail_table(
        document,
        ["항목", "결정"],
        [
            ("모델", "모델 A만 먼저 개발. 모델 B와 Unity 연결은 보류"),
            ("출력", "매초 센서값으로 정상·초기·중기·심각 4단계 판정"),
            ("데이터", "dataset/final_training_dataset_solver_latest.csv 완성본 사용"),
            ("알고리즘", "LightGBM 다중 분류"),
            ("검증", "cycle_id를 묶은 StratifiedGroupKFold 5겹 검증"),
        ],
    )

    document.add_heading("2. 파일별 역할", level=1)
    add_label_detail_table(
        document,
        ["파일", "역할"],
        [
            ("config.py", "Java properties처럼 경로·센서명·단계·검증 횟수 관리"),
            ("preparation.py", "기존 CSV 읽기, 무결성 검사, 입력과 정답 준비. 새 CSV는 만들지 않음"),
            ("train.py", "5겹 검증을 실행하고 전체 데이터로 최종 모델 학습"),
            ("evaluation.py", "severity_level 정답과 예측을 비교해 평가 결과 JSON 생성"),
        ],
    )

    document.add_heading("3. 모델 입력값 26개", level=1)
    add_input_table(
        document,
        [
            ("운전 문맥", "cycle_second", "현재 사이클이 시작된 뒤 지난 초"),
            ("운전 문맥", "cycle_phase", "하강·가압 유지·상승 등 현재 동작 단계"),
            ("운전 문맥", "target_chamber", "목표 압력을 가하는 실린더 방향(cap 또는 rod)"),
            ("온도", "temperature_c", "유압 오일 온도(°C)"),
            ("펌프", "pump_rpm", "펌프 회전수(RPM)"),
            ("밸브", "valve_cmd_pa", "P→A 유로의 명령 개도"),
            ("밸브", "valve_pos_pa", "P→A 유로의 실제 개도"),
            ("밸브", "valve_cmd_pb", "P→B 유로의 명령 개도"),
            ("밸브", "valve_pos_pb", "P→B 유로의 실제 개도"),
            ("밸브", "valve_cmd_at", "A→T 유로의 명령 개도"),
            ("밸브", "valve_pos_at", "A→T 유로의 실제 개도"),
            ("밸브", "valve_cmd_bt", "B→T 유로의 명령 개도"),
            ("밸브", "valve_pos_bt", "B→T 유로의 실제 개도"),
            ("유량", "flow_rate_L_min", "펌프 공급 유량(L/min)"),
            ("유량", "return_flow_L_min", "탱크로 돌아가는 리턴 유량(L/min)"),
            ("압력", "pump_in_pressure_bar", "펌프 입구 압력(bar)"),
            ("압력", "pump_out_pressure_bar", "펌프 출구 압력(bar)"),
            ("압력", "cyl_cap_pressure_bar", "실린더 cap 측 압력(bar)"),
            ("압력", "cyl_rod_pressure_bar", "실린더 rod 측 압력(bar)"),
            ("계산값", "pump_delta_pressure_bar", "펌프 출구와 입구의 압력 차(bar)"),
            ("계산값", "pump_head_m", "펌프가 만드는 환산 양정(m)"),
            ("계산값", "hydraulic_power_kW", "압력과 유량으로 계산한 유압 동력(kW)"),
            ("계산값", "max_velocity_m_s", "계통에서 계산된 최대 유속(m/s)"),
            ("계산값", "target_pressure_error_bar", "목표 압력과 실제 압력의 차(bar)"),
            ("운전 설정", "target_pressure_bar", "현재 동작의 목표 압력(bar)"),
            ("운전 설정", "relief_set_pressure_bar", "릴리프 밸브 설정 압력(bar)"),
        ],
    )

    document.add_heading("4. 정답·그룹·제외값", level=1)
    add_label_paragraph(
        document,
        "정답",
        "severity_level 컬럼. none=정상, initial=초기, moderate=중기, severe=심각",
    )
    add_label_paragraph(
        document,
        "그룹",
        "cycle_id는 모델 입력이 아니며 같은 사이클이 학습과 검증에 섞이지 않게 나눌 때만 사용",
    )
    add_label_paragraph(
        document,
        "제외",
        "모든 gt_*와 fault_class, fault_severity, onset_type, is_anomaly. 정답 유출을 막기 위해 모델 입력에 사용하지 않음",
    )

    document.add_heading("5. 5겹 검증 방식", level=1)
    add_label_paragraph(
        document,
        "그룹 분리",
        "같은 cycle_id에 속한 여러 초의 행은 반드시 함께 이동하므로 학습과 검증에 동시에 들어가지 않음",
    )
    add_label_paragraph(
        document,
        "저장 결과",
        "5겹 accuracy와 macro F1, 단계별 precision·recall·F1, 평균·표준편차, 혼동행렬, cycle_second별 정확도",
    )

    document.add_heading("6. 모델 결과값", level=1)
    add_label_paragraph(
        document,
        "한 행의 판정",
        "none·initial·moderate·severe 중 하나와 한글 단계(정상·초기·중기·심각)",
    )
    add_label_paragraph(
        document,
        "단계별 확률",
        "네 단계 각각의 확률과 그중 가장 큰 값인 신뢰도. 첫 구현에서는 별도 예측 화면 없이 모델 정보로 저장",
    )

    document.add_heading("7. 검증 결과값", level=1)
    add_metric_table(
        document,
        [
            ("accuracy", "전체 검증 행 중 열화 단계를 정확히 맞힌 비율"),
            ("macro_f1", "정상·초기·중기·심각을 같은 비중으로 계산한 F1 평균"),
            ("weighted_f1", "각 단계의 데이터 수를 반영한 F1 평균"),
            ("precision", "특정 단계라고 예측한 것 중 실제로 그 단계인 비율"),
            ("recall", "실제 특정 단계 데이터 중 모델이 찾아낸 비율"),
            ("f1", "precision과 recall의 균형 점수"),
            ("support", "평가에 사용된 단계별 행 수"),
            ("confusion_matrix", "실제 단계와 예측 단계가 어떻게 엇갈렸는지 보여주는 교차표"),
            ("fold_mean / fold_std", "5겹 점수의 평균과 분할에 따른 흔들림 정도"),
            ("accuracy_by_cycle_second", "사이클 시작 후 몇 초인지에 따른 매초 판정 정확도"),
        ],
    )
    add_label_paragraph(
        document,
        "무결성 결과",
        "data_validation_result.json에 행·사이클 수, 결측값, 무한대, 컬럼 누락, 단계 충돌, 통과 여부 저장",
    )
    add_label_paragraph(
        document,
        "평가 결과",
        "evaluation_results.json에 5겹 점수, 단계별 지표, 혼동행렬, 초별 정확도와 검증 범위 저장",
    )
    add_label_paragraph(
        document,
        "최종 모델",
        "degradation_stage_lightgbm.joblib에 모델, 26개 입력 순서, 단계 순서와 학습 정보 저장",
    )

    document.add_heading("8. 검증 결과를 읽는 방법", level=1)
    add_label_paragraph(
        document,
        "현재 의미",
        "솔버와 시나리오로 생성한 데이터 안에서 네 단계를 얼마나 잘 구분하는지 확인하는 내부 검증",
    )
    add_label_paragraph(
        document,
        "현재 한계",
        "실제 공장 정확도를 증명하는 결과는 아님. 실제 센서 이력과 점검으로 확인된 단계가 있는 공정 데이터로 외부 검증이 필요",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


if __name__ == "__main__":
    print(build_document())
