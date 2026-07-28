"""사용자가 빠르게 읽을 수 있는 예지보전 요약 DOCX를 생성한다.

이 문서는 발표 자료가 아니라 현재 프로젝트 의사결정을 위한 짧은 참고용 문서다.
기존 문서를 수정하지 않고 별도 파일만 생성한다.
"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "DigitalTwin_예지보전_핵심요약.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
HEADER_FILL = "E8EEF5"
TEXT = "1F1F1F"
TABLE_WIDTH_DXA = 9360


def _set_font(run, size: float, bold: bool = False, color: str = TEXT) -> None:
    """한글이 Word에서 빠지지 않도록 한글/영문 글꼴을 모두 명시한다."""
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def _set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_widths(table, widths: list[int]) -> None:
    """표 너비를 고정해 Word와 PDF 렌더링에서 줄바꿈이 달라지지 않게 한다."""
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_w = cell._tc.tcPr.tcW
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    _set_font(run, 16 if level == 1 else 12.5, bold=True, color=BLUE if level == 1 else DARK_BLUE)


def _add_body(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.2
    if bold_prefix and text.startswith(bold_prefix):
        _set_font(paragraph.add_run(bold_prefix), 10.5, bold=True)
        _set_font(paragraph.add_run(text[len(bold_prefix):]), 10.5)
    else:
        _set_font(paragraph.add_run(text), 10.5)


def _add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    _set_table_widths(table, widths)
    for cell, text in zip(table.rows[0].cells, headers):
        _shade(cell, HEADER_FILL)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(paragraph.add_run(text), 9.3, bold=True, color=DARK_BLUE)
    for values in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, values):
            paragraph = cell.paragraphs[0]
            _set_font(paragraph.add_run(text), 9.1)
    _set_table_widths(table, widths)


def build_document(output_path: Path = OUTPUT_PATH) -> Path:
    """현재 대화에서 합의한 핵심만 담은 짧은 참고 문서를 만든다."""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.2

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    _set_font(title.add_run("DigitalTwin 예지보전 핵심 요약"), 20, bold=True, color=DARK_BLUE)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(10)
    _set_font(subtitle.add_run("현재 데이터 기준 | 2026-07-21"), 9.5, color="666666")

    _add_heading(doc, "1. 만들고 싶은 예지보전", 1)
    _add_table(
        doc,
        ["결과", "의미"],
        [
            ["열화 단계", "정상 → 초기 → 중기 → 심각 → 고장"],
            ["열화 속도", "건강도가 몇 사이클 동안 얼마나 나빠지는지"],
            ["고장 시점 / RUL", "언제 고장나는지, 몇 시간·사이클 남았는지"],
            ["부품 위험도", "펌프·배관·밸브·실린더·오일계통 중 점검 우선순위"],
        ],
        [2500, 6860],
    )

    _add_heading(doc, "2. 모델 A와 B", 1)
    _add_table(
        doc,
        ["모델", "역할", "현재 데이터 판단"],
        [
            ["A. 건강·열화", "건강도, 열화 단계, 속도, 고장 시점·RUL", "현재 상태 4단계만 가능"],
            ["B. 부품 위험도", "위험 부품 계통의 확률", "현재 데이터로 가능"],
        ],
        [1800, 4200, 3360],
    )

    _add_heading(doc, "3. 현재 데이터로 가능한 것과 불가능한 것", 1)
    _add_table(
        doc,
        ["기능", "판정", "이유"],
        [
            ["정상·초기·중기·심각 판정", "가능", "severity_level 4개 라벨 존재"],
            ["5단계 고장 판정", "불가", "failure 라벨 없음"],
            ["열화 속도", "불가", "동일 설비의 장기 연속 이력 없음"],
            ["고장 시점·RUL", "불가", "run_id·누적 수명·실제 고장 시점 없음"],
            ["부품 위험도", "가능", "fault_class를 부품 계통 라벨로 묶어 사용"],
        ],
        [3000, 1500, 4860],
    )

    _add_heading(doc, "4. 모델 A에 추가로 필요한 데이터", 1)
    _add_table(
        doc,
        ["필수 데이터", "용도"],
        [
            ["equipment_id / run_id", "같은 설비의 정비 후 → 고장 전 연속 운전 묶음"],
            ["lifecycle_cycle / timestamp", "누적 수명 축과 실제 시간"],
            ["health_index", "예: 100=정상, 0=고장인 건강도 정답"],
            ["health_stage 5개", "normal, early, middle, severe, failure"],
            ["failure_cycle / failure_time", "고장 시점과 RUL 계산 기준"],
            ["실제 센서 이력", "압력·유량·온도·진동·전류 등 시간순 값"],
        ],
        [3300, 6060],
    )

    _add_heading(doc, "5. 권장 순서", 1)
    _add_body(doc, "1) 최신 솔버 결과로 데이터 정합성을 맞춘 뒤 모델 B를 먼저 학습합니다.")
    _add_body(doc, "2) 설비별 정상→고장 연속 이력을 생성한 뒤 모델 A의 5단계·열화 속도 기능을 만듭니다.")
    _add_body(doc, "3) 고장 시점과 RUL은 연속 이력이 충분히 쌓인 뒤 모델 A에 추가합니다.")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path


if __name__ == "__main__":
    print(build_document())
