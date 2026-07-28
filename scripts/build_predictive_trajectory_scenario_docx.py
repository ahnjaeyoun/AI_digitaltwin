"""간단한 팀 공유용 예지보전 데이터 시나리오 DOCX를 생성한다."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "DigitalTwin_예지보전_열화궤적_데이터시나리오_요약.docx"


def set_korean_font(run, size: float = 11, bold: bool = False) -> None:
    """한글이 깨지지 않도록 글꼴을 명시한다."""
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:ascii"), "맑은 고딕")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "맑은 고딕")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(size)
    run.bold = bold


def configure_styles(document: Document) -> None:
    """장식을 줄이고 기본적인 읽기 편한 형식만 설정한다."""
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)

    normal = document.styles["Normal"]
    normal.font.name = "맑은 고딕"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15

    heading1 = document.styles["Heading 1"]
    heading1.font.name = "맑은 고딕"
    heading1._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    heading1.font.size = Pt(13)
    heading1.font.bold = True
    heading1.font.color.rgb = RGBColor(0, 0, 0)
    heading1.paragraph_format.space_before = Pt(10)
    heading1.paragraph_format.space_after = Pt(4)

    bullet = document.styles["List Bullet"]
    bullet.font.name = "맑은 고딕"
    bullet._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    bullet.font.size = Pt(10.5)
    bullet.paragraph_format.left_indent = Inches(0.38)
    bullet.paragraph_format.first_line_indent = Inches(-0.19)
    bullet.paragraph_format.space_after = Pt(2)
    bullet.paragraph_format.line_spacing = 1.15


def add_bullet(document: Document, text: str) -> None:
    """실제 Word 글머리표 문단을 추가한다."""
    paragraph = document.add_paragraph(style="List Bullet")
    set_korean_font(paragraph.add_run(text), size=10.5)


def add_labeled_line(document: Document, label: str, content: str) -> None:
    """짧은 항목 이름과 설명을 한 줄에 작성한다."""
    paragraph = document.add_paragraph()
    set_korean_font(paragraph.add_run(f"{label}: "), size=10.5, bold=True)
    set_korean_font(paragraph.add_run(content), size=10.5)


def build_document() -> None:
    """합의한 시나리오만 남긴 짧은 DOCX를 작성한다."""
    document = Document()
    configure_styles(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(10)
    set_korean_font(
        title.add_run("DigitalTwin 예지보전 데이터 시나리오"),
        size=17,
        bold=True,
    )

    document.add_heading("1. 목표", level=1)
    document.add_paragraph(
        "연속된 센서 변화로 예상 고장 시점, 남은 사이클(RUL), 남은 시간과 열화 속도를 예측합니다. "
        "고장 유형과 위험 부품 판정은 다른 팀원의 모델에서 담당합니다."
    )

    document.add_heading("2. 운전 단위", level=1)
    add_bullet(document, "1사이클은 상승 6초 + 유지 3초 + 하강 6초로 구성합니다.")
    add_bullet(document, "초 단위 센서값과 구간별·사이클별 대표값을 함께 저장합니다.")

    document.add_heading("3. 열화 이력 시나리오", level=1)
    document.add_paragraph(
        "정비 완료 → 정상 → 초기 열화 → 중기 열화 → 심각 열화 → 고장/안전 한계 도달 → 정지 → 정비 → 새 이력 시작"
    )
    add_bullet(document, "trajectory_id 하나는 정비 후 정상 운전부터 다음 고장까지의 연속 이력입니다.")
    add_bullet(document, "고장 후에는 운전을 계속하지 않고 정비 후 새 trajectory_id로 다시 시작합니다.")
    add_bullet(document, "health_stage는 정상·초기·중기·심각의 4단계이며, 고장은 별도의 종료 이벤트입니다.")

    document.add_heading("4. 관리도와 z-score", level=1)
    add_bullet(document, "관리도는 AI 출력이 아니라 열화 단계와 failure_cycle을 정하는 기준입니다.")
    add_bullet(document, "상승·유지·하강 구간별 정상 평균과 표준편차를 따로 계산합니다.")
    add_bullet(document, "z-score는 현재 센서값이 정상 기준에서 얼마나 벗어났는지 나타내는 내부 계산값입니다.")
    add_bullet(document, "z-score를 반드시 AI 입력·출력 컬럼으로 사용할 필요는 없습니다.")

    document.add_heading("5. 점진 열화와 급성 고장", level=1)
    add_labeled_line(
        document,
        "점진 열화",
        "예지보전의 핵심 학습 대상입니다. 여러 사이클에 걸친 센서 추세로 고장 시점과 RUL을 예측합니다.",
    )
    add_labeled_line(
        document,
        "급성 고장",
        "사전 징후가 있으면 빠른 열화로 학습하고, 징후가 전혀 없으면 RUL 학습에서 제외하여 이상 감지용으로 사용합니다.",
    )

    document.add_heading("6. 핵심 컬럼", level=1)
    add_bullet(document, "이력: trajectory_id, cycle_id, cycle_second, phase, timestamp")
    add_bullet(document, "상태: health_stage, degradation_mode, machine_status")
    add_bullet(document, "시점: onset_cycle, failure_cycle, cycles_to_failure, time_to_failure_hours")
    add_bullet(document, "원인: fault_class, fault_component, maintenance_action")
    add_bullet(document, "입력: 솔버에서 계산된 실제 센서값과 구간별·사이클별 대표값")

    document.add_heading("7. 데이터 생성 방법", level=1)
    add_bullet(document, "고장 유형마다 점진 열화 이력을 우선 100~200개 생성합니다.")
    add_bullet(document, "정상 운전 조건, 열화 시작 시점, 열화 속도, 센서 노이즈와 고장 시점을 바꿉니다.")
    add_bullet(document, "단계를 임의로 붙이지 않고 솔버 출력과 관리도 기준으로 결정합니다.")
    add_bullet(document, "현재 데이터는 사이클 간 연속 이력이 부족하므로 이 구조로 재생성이 필요합니다.")

    document.save(OUTPUT_PATH)
    print(f"생성 완료: {OUTPUT_PATH}")


def verify_document() -> None:
    """저장된 문서가 열리고 핵심 내용이 있는지 확인한다."""
    verified = Document(OUTPUT_PATH)
    all_text = "\n".join(paragraph.text for paragraph in verified.paragraphs)
    required = [
        "상승 6초 + 유지 3초 + 하강 6초",
        "정상·초기·중기·심각의 4단계",
        "관리도는 AI 출력이 아니라",
        "점진 열화",
        "급성 고장",
        "cycles_to_failure",
    ]
    missing = [text for text in required if text not in all_text]
    if missing:
        raise RuntimeError(f"핵심 내용 누락: {missing}")
    if verified.tables:
        raise RuntimeError("간단 문서에 불필요한 표가 남아 있습니다.")
    print(f"구조 검증 완료: 문단 {len(verified.paragraphs)}개, 표 0개")


if __name__ == "__main__":
    build_document()
    verify_document()
